"""文件解析器 — 支持 md/txt/docx/xlsx/csv，统一返回纯文本。

支持插件化注册：通过 register_parser() 可注册自定义解析器。
"""

import logging
from pathlib import Path
from typing import Callable, Dict

logger = logging.getLogger(__name__)

ParserFunc = Callable[[Path], str]


_PARSERS: Dict[str, ParserFunc] = {}


def register_parser(suffix: str, func: ParserFunc) -> None:
    """注册自定义文件解析器。

    Args:
        suffix: 文件后缀（如 '.pdf'，不区分大小写）
        func: 解析函数，接收 Path 返回纯文本
    """
    _PARSERS[suffix.lower()] = func
    logger.info("注册解析器: %s -> %s", suffix, func.__name__)


def parse_file(path: Path) -> str:
    """根据文件后缀自动选择解析器，返回纯文本。"""
    suffix = path.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"Unsupported file type: {suffix}")
    return parser(path)


def _parse_md(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _parse_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _parse_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"# {sheet.title}\n\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def _parse_csv(path: Path) -> str:
    import csv
    text = path.read_text(encoding="utf-8")
    reader = csv.reader(text.splitlines())
    rows = []
    for row in reader:
        if any(cell.strip() for cell in row):
            rows.append(" | ".join(row))
    return "\n".join(rows)


# 注册内置解析器
for _ext, _func in [
    (".md", _parse_md),
    (".markdown", _parse_md),
    (".txt", _parse_txt),
    (".docx", _parse_docx),
    (".xlsx", _parse_xlsx),
    (".csv", _parse_csv),
]:
    register_parser(_ext, _func)

SUPPORTED_EXTENSIONS = list(_PARSERS.keys())

